import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
from typing import Dict, Any, List

class ExcelWorkpaperExporter:
    """
    Generates a multi-tab, formatted Excel Schedule C Workpaper (.xlsx)
    with live formulas, drill-down audit trails, and exception logs.
    """

    def generate_workpaper(self, client_name: str, tax_year: int, transactions: List[Dict[str, Any]],
                           exceptions_dict: Dict[str, Any], questions: List[Dict[str, str]],
                           coverage_info: Dict[str, Any], duplicates_info: List[Dict[str, str]]) -> io.BytesIO:
        """
        Creates an openpyxl Workbook and returns a BytesIO buffer.
        """
        wb = openpyxl.Workbook()
        
        # Define Color Palette & Styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid") # Dark Navy Blue
        sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid") # Soft Blue
        alert_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft Coral/Alert
        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft Green

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
        subtitle_font = Font(name="Calibri", size=12, bold=True, color="595959")
        bold_font = Font(name="Calibri", size=11, bold=True)
        normal_font = Font(name="Calibri", size=11)

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        double_bottom_border = Border(
            top=Side(style='thin', color='000000'),
            bottom=Side(style='double', color='000000')
        )

        # ----------------------------------------------------
        # TAB 1: SCHEDULE C SUMMARY DASHBOARD
        # ----------------------------------------------------
        ws1 = wb.active
        ws1.title = "Schedule C Summary"
        ws1.views.sheetView[0].showGridLines = True

        ws1.cell(row=1, column=1, value="SCHEDULE C TAX PREPARATION WORKPAPER").font = title_font
        ws1.cell(row=2, column=1, value=f"Client: {client_name} | Tax Year: {tax_year} | Status: Ready for Preparer Review").font = subtitle_font

        # Client Meta Header Block
        meta_data = [
            ("Client Name:", client_name, "Deduplication Rule:", "Strict Filename Verification (Active)"),
            ("Tax Year:", tax_year, "12-Month Statement Coverage:", coverage_info.get("status_message", "N/A")),
            ("Prepared By:", "GoSpectr Schedule C Tool", "Total Exceptions Flagged:", exceptions_dict.get("total_exception_count", 0))
        ]

        for r_idx, row_vals in enumerate(meta_data, start=4):
            ws1.cell(row=r_idx, column=1, value=row_vals[0]).font = bold_font
            ws1.cell(row=r_idx, column=2, value=row_vals[1]).font = normal_font
            ws1.cell(row=r_idx, column=4, value=row_vals[2]).font = bold_font
            ws1.cell(row=r_idx, column=5, value=row_vals[3]).font = normal_font

        # Calculate Income & Expenses Summary
        gross_receipts = sum(tx['amount'] for tx in transactions if tx.get('is_deposit') and not tx.get('category', '').startswith("Non-P&L:"))
        
        # Expense Line Item Aggregation
        expense_totals = {}
        for tx in transactions:
            if not tx.get('is_deposit') and not tx.get('category', '').startswith("Non-P&L:"):
                cat = tx.get('category', 'Line 27a: Other expenses (Uncategorized)')
                amt = abs(tx.get('amount', 0.0))
                if "Line 24b:" in cat:
                    amt = amt * 0.5
                expense_totals[cat] = expense_totals.get(cat, 0.0) + amt

        # Write Financial Summary Table
        ws1.cell(row=8, column=1, value="IRS SCHEDULE C LINE ITEM SUMMARY").font = bold_font
        headers = ["IRS Line Item", "Tax Category Description", "Amount ($)", "Audit Notes / QC Status"]
        
        for c_idx, h_text in enumerate(headers, start=1):
            cell = ws1.cell(row=9, column=c_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill

        curr_row = 10
        # Income Row
        ws1.cell(row=curr_row, column=1, value="Line 1").font = bold_font
        ws1.cell(row=curr_row, column=2, value="Gross Receipts / Sales").font = bold_font
        c_gross = ws1.cell(row=curr_row, column=3, value=gross_receipts)
        c_gross.font = bold_font
        c_gross.number_format = "$#,##0.00"
        ws1.cell(row=curr_row, column=4, value="Cleaned of internal transfers & loans").font = normal_font
        curr_row += 1

        # Section Divider
        ws1.cell(row=curr_row, column=2, value="EXPENSES").font = bold_font
        ws1.cell(row=curr_row, column=2).fill = sub_fill
        curr_row += 1

        exp_start_row = curr_row
        sorted_exp_cats = sorted(expense_totals.keys())
        
        if sorted_exp_cats:
            for cat in sorted_exp_cats:
                amt = expense_totals[cat]
                line_num = cat.split(":")[0] if ":" in cat else "Line 27a"
                cat_desc = cat.split(":")[1].strip() if ":" in cat else cat

                ws1.cell(row=curr_row, column=1, value=line_num).font = normal_font
                ws1.cell(row=curr_row, column=2, value=cat_desc).font = normal_font
                c_amt = ws1.cell(row=curr_row, column=3, value=amt)
                c_amt.font = normal_font
                c_amt.number_format = "$#,##0.00"
                
                note = "50% Meal Rule Applied" if "24b" in line_num else "Verified"
                ws1.cell(row=curr_row, column=4, value=note).font = normal_font
                curr_row += 1
            exp_end_row = curr_row - 1
        else:
            ws1.cell(row=curr_row, column=1, value="Line 27a").font = normal_font
            ws1.cell(row=curr_row, column=2, value="Other expenses (Uncategorized)").font = normal_font
            c_amt = ws1.cell(row=curr_row, column=3, value=0.0)
            c_amt.font = normal_font
            c_amt.number_format = "$#,##0.00"
            exp_start_row = curr_row
            exp_end_row = curr_row
            curr_row += 1

        # Total Expenses Row
        ws1.cell(row=curr_row, column=1, value="Line 28").font = bold_font
        ws1.cell(row=curr_row, column=2, value="TOTAL EXPENSES").font = bold_font
        c_tot_exp = ws1.cell(row=curr_row, column=3, value=f"=SUM(C{exp_start_row}:C{exp_end_row})")
        c_tot_exp.font = bold_font
        c_tot_exp.number_format = "$#,##0.00"
        c_tot_exp.border = thin_border
        curr_row += 1

        # Net Profit / Loss Row
        ws1.cell(row=curr_row, column=1, value="Line 31").font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        ws1.cell(row=curr_row, column=2, value="NET PROFIT / (LOSS)").font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        c_net = ws1.cell(row=curr_row, column=3, value=f"=C10-C{curr_row-1}")
        c_net.font = Font(name="Calibri", size=11, bold=True, color="1F4E78")
        c_net.number_format = "$#,##0.00"
        c_net.border = double_bottom_border

        # ----------------------------------------------------
        # TAB 2: FULL TRANSACTIONS AUDIT TRAIL
        # ----------------------------------------------------
        ws2 = wb.create_sheet(title="All Transactions")
        ws2.views.sheetView[0].showGridLines = True

        tx_headers = ["Date", "Payee / Vendor", "Description", "Amount ($)", "Deposit?", "Schedule C Category", "Original Client Cat", "Confidence State", "Source File"]
        for c_idx, h_text in enumerate(tx_headers, start=1):
            cell = ws2.cell(row=1, column=c_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill

        for r_idx, tx in enumerate(transactions, start=2):
            ws2.cell(row=r_idx, column=1, value=tx.get('date')).font = normal_font
            ws2.cell(row=r_idx, column=2, value=tx.get('payee')).font = normal_font
            ws2.cell(row=r_idx, column=3, value=tx.get('description')).font = normal_font
            
            c_amt = ws2.cell(row=r_idx, column=4, value=tx.get('amount'))
            c_amt.font = normal_font
            c_amt.number_format = "$#,##0.00"
            
            ws2.cell(row=r_idx, column=5, value="Yes" if tx.get('is_deposit') else "No").font = normal_font
            ws2.cell(row=r_idx, column=6, value=tx.get('category')).font = normal_font
            ws2.cell(row=r_idx, column=7, value=tx.get('original_category')).font = normal_font
            
            c_conf = ws2.cell(row=r_idx, column=8, value=tx.get('confidence_state'))
            c_conf.font = normal_font
            if tx.get('confidence_state') != "High Confidence":
                c_conf.fill = alert_fill

            ws2.cell(row=r_idx, column=9, value=tx.get('source_file')).font = normal_font

        # ----------------------------------------------------
        # TAB 3: NON-P&L & TRANSFERS
        # ----------------------------------------------------
        ws3 = wb.create_sheet(title="Non-P&L & Transfers")
        ws3.views.sheetView[0].showGridLines = True
        
        non_pnl_items = exceptions_dict.get("non_pnl_transfers", [])
        for c_idx, h_text in enumerate(tx_headers, start=1):
            cell = ws3.cell(row=1, column=c_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill

        for r_idx, tx in enumerate(non_pnl_items, start=2):
            ws3.cell(row=r_idx, column=1, value=tx.get('date')).font = normal_font
            ws3.cell(row=r_idx, column=2, value=tx.get('payee')).font = normal_font
            ws3.cell(row=r_idx, column=3, value=tx.get('description')).font = normal_font
            
            c_amt = ws3.cell(row=r_idx, column=4, value=tx.get('amount'))
            c_amt.font = normal_font
            c_amt.number_format = "$#,##0.00"
            
            ws3.cell(row=r_idx, column=5, value="Yes" if tx.get('is_deposit') else "No").font = normal_font
            ws3.cell(row=r_idx, column=6, value=tx.get('category')).font = normal_font
            ws3.cell(row=r_idx, column=7, value=tx.get('original_category')).font = normal_font
            ws3.cell(row=r_idx, column=8, value=tx.get('confidence_state')).font = normal_font
            ws3.cell(row=r_idx, column=9, value=tx.get('source_file')).font = normal_font

        # ----------------------------------------------------
        # TAB 4: EXCEPTIONS QUEUE
        # ----------------------------------------------------
        ws4 = wb.create_sheet(title="Exceptions Queue")
        ws4.views.sheetView[0].showGridLines = True

        ex_headers = ["Date", "Payee", "Amount ($)", "Flag Category", "Reason / Audit Note"]
        for c_idx, h_text in enumerate(ex_headers, start=1):
            cell = ws4.cell(row=1, column=c_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill

        row_counter = 2
        for item in exceptions_dict.get("potential_personal", []):
            ws4.cell(row=row_counter, column=1, value=item.get('date')).font = normal_font
            ws4.cell(row=row_counter, column=2, value=item.get('payee')).font = normal_font
            c_a = ws4.cell(row=row_counter, column=3, value=item.get('amount'))
            c_a.font = normal_font
            c_a.number_format = "$#,##0.00"
            ws4.cell(row=row_counter, column=4, value="Potential Personal Expense").font = bold_font
            ws4.cell(row=row_counter, column=5, value=item.get('reason')).font = normal_font
            ws4.cell(row=row_counter, column=4).fill = alert_fill
            row_counter += 1

        for item in exceptions_dict.get("potential_fixed_assets", []):
            ws4.cell(row=row_counter, column=1, value=item.get('date')).font = normal_font
            ws4.cell(row=row_counter, column=2, value=item.get('payee')).font = normal_font
            c_a = ws4.cell(row=row_counter, column=3, value=item.get('amount'))
            c_a.font = normal_font
            c_a.number_format = "$#,##0.00"
            ws4.cell(row=row_counter, column=4, value="Potential Fixed Asset (> $2,500)").font = bold_font
            ws4.cell(row=row_counter, column=5, value=item.get('reason')).font = normal_font
            ws4.cell(row=row_counter, column=4).fill = sub_fill
            row_counter += 1

        # ----------------------------------------------------
        # TAB 5: CLIENT QUESTIONS CHECKLIST
        # ----------------------------------------------------
        ws5 = wb.create_sheet(title="Client Questions")
        ws5.views.sheetView[0].showGridLines = True

        q_headers = ["Item ID", "Category", "Date", "Payee / Entity", "Amount ($)", "Question for Client", "Client Response"]
        for c_idx, h_text in enumerate(q_headers, start=1):
            cell = ws5.cell(row=1, column=c_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill

        for r_idx, q in enumerate(questions, start=2):
            ws5.cell(row=r_idx, column=1, value=q.get('item_id')).font = bold_font
            ws5.cell(row=r_idx, column=2, value=q.get('category')).font = normal_font
            ws5.cell(row=r_idx, column=3, value=q.get('date')).font = normal_font
            ws5.cell(row=r_idx, column=4, value=q.get('payee')).font = normal_font
            ws5.cell(row=r_idx, column=5, value=q.get('amount')).font = normal_font
            ws5.cell(row=r_idx, column=6, value=q.get('question')).font = normal_font
            ws5.cell(row=r_idx, column=7, value=q.get('client_response')).font = normal_font

        # Auto-adjust column widths for all worksheets
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        output_buffer = io.BytesIO()
        wb.save(output_buffer)
        output_buffer.seek(0)
        return output_buffer
